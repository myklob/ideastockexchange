"""Tests for the scoring engine. Run: python3 -m unittest discover -s tools/static-site -p 'test_*.py' -v

These pin the rules, not the numbers, so they stay meaningful as content changes. Where a number is pinned it is
derived in the test from the rule it checks, so a deliberate rule change fails loudly and a content edit does not.
The CI workflow runs this before it will publish anything.
"""
import os, sys, unittest
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from score_reference import Model, CONSTS
from confidence import Confidence

K, UNARG, DEFLINK, DEFIMP, DEFUNIQ = (CONSTS[k] for k in ('K', 'UNARG', 'DEFLINK', 'DEFIMP', 'DEFUNIQ'))
ENTRY = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'ISE_Data_Entry.xlsx')


def corpus(rows, kind='belief'):
    """A parent page with `rows` = [(truth_of_child, side)], each child a page argued to that truth.

    A child is given its truth directly by attaching enough neutral-free structure is not possible, so instead
    each child is a leaf whose truth we set by construction: we stub the model's truth for that page.
    """
    pages = [{'id': 1, 'kind': kind, 'text': 'parent'}]
    edges = []
    for i, (t, side) in enumerate(rows, start=2):
        pages.append({'id': i, 'kind': 'claim', 'text': f'child {i}'})
        edges.append({'id': i, 'page_id': 1, 'section': 'argument', 'side': side, 'position': i, 'claim_id': i})
    m = Model(pages, CONSTS, edges=edges)
    truths = {i: t for i, (t, _) in enumerate(rows, start=2)}
    real = m.evaluate
    m.evaluate = lambda pid, _seen=(): ({'truth': truths[pid], 'belief': 0.0, 'pro': 0, 'con': 0, 'supp': 0,
                                         'weak': 0, 'pred': 0, 'impact': None, 'raw': None}
                                        if pid in truths else real(pid, _seen))
    m.conf = lambda pid: 1.0
    return m


class TestRowContribution(unittest.TestCase):
    """sign x (2 x Truth - 1) x Confidence x Link x Imp x Uniq."""

    def test_unargued_claim_contributes_exactly_zero(self):
        """The whole point of the signed scale: listing a claim is worth nothing until someone argues it."""
        m = corpus([(UNARG, 'agree')] * 10)
        self.assertEqual(m.evaluate(1)['belief'], 0.0)
        self.assertEqual(m.evaluate(1)['truth'], 0.5)

    def test_refuted_reason_to_agree_counts_against(self):
        m = corpus([(0.1, 'agree')])
        self.assertLess(m.evaluate(1)['belief'], 0)

    def test_refuted_objection_counts_as_support(self):
        """A disagree row argued false is evidence FOR the belief. Weight lands on the side its sign puts it on."""
        m = corpus([(0.1, 'disagree')])
        self.assertGreater(m.evaluate(1)['belief'], 0)
        self.assertGreater(m.evaluate(1)['pro'], 0)

    def test_proven_reason_contributes_its_full_multiplier_product(self):
        m = corpus([(1.0, 'agree')])
        self.assertAlmostEqual(m.evaluate(1)['belief'], 1.0 * DEFLINK * DEFIMP * DEFUNIQ)

    def test_contribution_is_bounded_by_one_in_magnitude(self):
        for t in (0.0, 0.25, 0.5, 0.75, 1.0):
            for side in ('agree', 'disagree'):
                self.assertLessEqual(abs(corpus([(t, side)]).evaluate(1)['belief']), 1.0)


class TestPaddingExploit(unittest.TestCase):
    """Regression: before the signed scale, padding a page with claims known to be false RAISED its score."""

    def test_padding_with_refuted_arguments_lowers_the_score(self):
        base = corpus([(0.9, 'agree'), (0.8, 'agree')])
        padded = corpus([(0.9, 'agree'), (0.8, 'agree')] + [(0.1, 'agree')] * 10)
        self.assertLess(padded.evaluate(1)['belief'], base.evaluate(1)['belief'])
        self.assertLess(padded.evaluate(1)['truth'], base.evaluate(1)['truth'])

    def test_padding_with_unargued_claims_changes_nothing(self):
        base = corpus([(0.9, 'agree')])
        padded = corpus([(0.9, 'agree')] + [(UNARG, 'agree')] * 50)
        self.assertAlmostEqual(padded.evaluate(1)['belief'], base.evaluate(1)['belief'])
        self.assertAlmostEqual(padded.evaluate(1)['truth'], base.evaluate(1)['truth'])


class TestInvariants(unittest.TestCase):
    def test_truth_always_within_zero_and_one(self):
        for rows in ([(1.0, 'agree')] * 30, [(0.0, 'agree')] * 30, [(1.0, 'agree'), (1.0, 'disagree')], []):
            t = corpus(rows).evaluate(1)['truth']
            self.assertGreaterEqual(t, 0.0); self.assertLessEqual(t, 1.0)

    def test_empty_page_reads_the_neutral_prior(self):
        self.assertAlmostEqual(corpus([]).evaluate(1)['truth'], 0.5)

    def test_balanced_evidence_reads_neutral(self):
        self.assertAlmostEqual(corpus([(1.0, 'agree'), (1.0, 'disagree')]).evaluate(1)['truth'], 0.5)

    def test_belief_equals_weight_for_minus_weight_against(self):
        m = corpus([(0.9, 'agree'), (0.2, 'agree'), (0.8, 'disagree')])
        e = m.evaluate(1)
        self.assertAlmostEqual(e['belief'], e['pos'] - e['neg'])


class TestRealCorpus(unittest.TestCase):
    """Integration against the shipped workbook. Skipped if openpyxl or the workbook is unavailable."""

    @classmethod
    def setUpClass(cls):
        try:
            import openpyxl  # noqa: F401
            from render_site import Corpus
        except ImportError as e:
            raise unittest.SkipTest(f'openpyxl missing: {e}')
        if not os.path.exists(ENTRY): raise unittest.SkipTest('workbook not present')
        cls.c = Corpus(ENTRY, 'test')

    def test_site_and_engine_agree_on_every_page(self):
        """Two independent implementations of the same rules. They must not drift."""
        for p in self.c.specs:
            s, e = self.c.stats(p), self.c.model.evaluate(p)
            for k in ('truth', 'belief', 'pro', 'con'):
                self.assertAlmostEqual(s[k], e[k], places=12, msg=f'page {p} field {k}')

    def test_every_truth_in_range(self):
        for p in self.c.specs:
            t = self.c.truth(p)
            self.assertGreaterEqual(t, 0.0, p); self.assertLessEqual(t, 1.0, p)

    def test_every_confidence_in_range(self):
        for p in self.c.specs:
            k = self.c.conf.of(p)
            self.assertGreaterEqual(k, 0.0, p); self.assertLessEqual(k, 1.0, p)

    def test_no_page_is_its_own_ancestor(self):
        """A cycle would make the recursive score meaningless, and would hang a naive scorer."""
        def walk(pid, seen):
            self.assertNotIn(pid, seen, f'cycle through page {pid}')
            sp = self.c.specs[pid]
            for key, side in (('args', 'agree'), ('args', 'disagree'), ('evid', 'for'), ('evid', 'against')):
                for d in sp.get(key, {}).get(side, []):
                    k = d.get('id')
                    if isinstance(k, int) and k in self.c.specs: walk(k, seen | {pid})
        for p in self.c.beliefs: walk(p, frozenset())

    def test_every_reference_resolves(self):
        """No row may point at a page that does not exist: a dangling reference is a silent zero."""
        for p, sp in self.c.specs.items():
            for key in ('x', 'y', 'z', 'supports'):
                v = sp.get(key)
                if isinstance(v, int): self.assertIn(v, self.c.specs, f'page {p}.{key} -> {v}')

    def test_scoring_is_deterministic(self):
        from render_site import Corpus
        again = Corpus(ENTRY, 'test')
        for p in self.c.specs:
            self.assertAlmostEqual(self.c.truth(p), again.truth(p), places=12)
            self.assertAlmostEqual(self.c.conf.of(p), again.conf.of(p), places=12)


class TestConfidence(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        try:
            import openpyxl  # noqa: F401
            from render_site import Corpus
        except ImportError as e:
            raise unittest.SkipTest(f'openpyxl missing: {e}')
        if not os.path.exists(ENTRY): raise unittest.SkipTest('workbook not present')
        cls.c = Corpus(ENTRY, 'test')

    def test_a_page_with_no_work_behind_it_has_no_confidence(self):
        """Nothing argued, nothing listed, nothing cited: confidence is exactly 0, so the page moves nothing
        above it however often it is listed as a reason."""
        from confidence import Confidence, TableCorpus
        k = Confidence(TableCorpus([{'id': 1, 'kind': 'claim', 'text': 'a bare assertion'}], []))
        self.assertEqual(k.of(1), 0.0)

    def test_confidence_comes_only_from_work_actually_done(self):
        """The mirror of the rule above, over the real corpus: no page may have confidence without rows
        beneath it, interests listed on it, or a source cited for it."""
        import evidence as EV
        for p, sp in self.c.specs.items():
            if self.c.conf.of(p) <= 0: continue
            work = (sp.get('args', {}).get('agree') or sp.get('args', {}).get('disagree')
                    or sp.get('evid', {}).get('for') or sp.get('evid', {}).get('against')
                    or sp.get('pred_true') or sp.get('pred_false') or sp.get('interests')
                    or EV.prior(sp)['classified'])
            self.assertTrue(work, f'page {p} has confidence {self.c.conf.of(p)} with no work behind it')

    def test_confidence_never_calls_the_scorer(self):
        """It must measure work done, not what the work concluded, or it recurses through truth."""
        import confidence as mod
        with open(mod.__file__) as fh: src = fh.read()
        for forbidden in ('.truth(', 'evaluate('):
            self.assertNotIn(forbidden, src, f'confidence.py must not call {forbidden}')

    def test_components_are_all_within_range(self):
        for p in self.c.specs:
            for name, v in self.c.conf.parts(p)['components'].items():
                self.assertGreaterEqual(v, 0.0, f'{p}.{name}'); self.assertLessEqual(v, 1.0, f'{p}.{name}')


class TestWorkbookFormulasMatchTheEngine(unittest.TestCase):
    """The Excel workbook is a second implementation of the same rules. LibreOffice cannot be driven in every
    environment, so rather than recalculating a 270-sheet workbook these check the formula strings the builder
    emits. A silent divergence between the workbook and the website is the failure mode this guards against."""

    @classmethod
    def setUpClass(cls):
        try:
            from openpyxl import Workbook
            import build_pages as bp
            from build_subpages import SubPage
        except ImportError as e:
            raise unittest.SkipTest(f'openpyxl missing: {e}')
        cls.bp, cls.Workbook, cls.SubPage = bp, Workbook, SubPage

    def _build(self, spec, kind=None):
        wb = self.Workbook(); ws = wb.active
        pg = self.SubPage(ws, dict(spec, kind=kind)) if kind else self.bp.Page(ws, spec)
        pg.build()
        return ws, pg

    def test_belief_row_formula_is_signed_and_confidence_gated(self):
        ws, pg = self._build({'args': {'agree': [{'id': 7, 'text': 'a'}], 'disagree': [{'id': 8, 'text': 'b'}]}})
        r = pg.rows['args'][0]
        agree, disagree = ws[f'J{r}'].value, ws[f'U{r}'].value
        for f, sign in ((agree, ''), (disagree, '-')):
            self.assertIn('(2*', f, 'row must use the signed (2 x Truth - 1) scale')
            self.assertIn('$W$3', f, 'row must multiply by the child page confidence mirror')
        self.assertIn(',-(2*', disagree, 'the disagree side must carry sign -1')
        self.assertNotIn(',-(2*', agree, 'the agree side must carry sign +1')

    def test_specialized_page_row_formula_matches(self):
        ws, pg = self._build({'args': {'agree': [{'id': 7, 'text': 'a'}], 'disagree': []}}, kind='linkage')
        f = ws[f'I{pg.rows["args"][0]}'].value
        self.assertIn('(2*', f); self.assertIn('$W$3', f)

    def test_totals_partition_by_sign_rather_than_summing(self):
        """With signed rows a plain SUM would net them; weight for and against must be separated."""
        ws, pg = self._build({'args': {'agree': [{'id': 7, 'text': 'a'}], 'disagree': [{'id': 8, 'text': 'b'}]}})
        tot = pg.rows['argtot']
        self.assertIn('SUMIF', ws[f'J{tot}'].value)
        self.assertIn('SUMIF', ws[f'U{tot}'].value)

    def test_every_page_exposes_a_confidence_mirror(self):
        for kind in (None, 'linkage', 'importance', 'media'):
            ws, _ = self._build({'args': {'agree': [], 'disagree': []}}, kind=kind)
            self.assertTrue(str(ws['W3'].value or '').startswith('='), f'{kind} page must mirror confidence at W3')


if __name__ == '__main__':
    unittest.main(verbosity=2)


class TestTheScorerOnItsOwn(unittest.TestCase):
    """A Model built with nothing attached to it. Two paths use it that way and neither is exercised anywhere
    else: `python3 score_reference.py <file>`, which the module docstring documents, and build_example.py when
    it has no corpus to gate on. So the default that both of them read was pinned by nothing, and changing it
    left every check in this repository green."""

    def test_confidence_defaults_to_one_so_the_zoning_reference_still_reproduces(self):
        """1.0 is not a neutral choice of number, it is the choice that leaves the scorer exactly as it was
        before confidence existed. The zoning example and its Python, SQL and PHP twins have to keep producing
        the numbers they were transcribed with, and they were transcribed without this gate."""
        from score_reference import Model, CONSTS
        m = Model({'pages': [{'id': 1, 'kind': 'belief'}], 'edges': []}, CONSTS)
        self.assertEqual(m.conf(1), 1.0)
        self.assertEqual(m.conf(999), 1.0, 'the default has to answer for a page it has never seen')

    def test_the_command_line_scores_a_file_without_anything_attached(self):
        """The docstring says you can run this file on a JSON corpus and get every page's truth. Nothing ran
        it, so the whole unattached path, default confidence included, was untested."""
        import json, os, subprocess, sys, tempfile
        here = os.path.dirname(os.path.abspath(__file__))
        corpus = {'constants': [{'name': k, 'value': v} for k, v in
                                __import__('score_reference').CONSTS.items()],
                  'pages': [{'id': 1, 'kind': 'belief', 'text': 'The claim under test'},
                            {'id': 2, 'kind': 'claim', 'text': 'A cited reason', 'etype': 'statistics'}],
                  'edges': [{'id': 1, 'page_id': 1, 'section': 'argument', 'side': 'agree',
                             'position': 1, 'claim_id': 2}]}
        d = tempfile.mkdtemp()
        path = os.path.join(d, 'corpus.json')
        with open(path, 'w') as fh: json.dump(corpus, fh)
        r = subprocess.run([sys.executable, os.path.join(here, 'score_reference.py'), path],
                           capture_output=True, text=True, timeout=60)
        self.assertEqual(r.returncode, 0, r.stderr[-400:])
        self.assertIn('0.95', r.stdout, 'the cited claim did not score where its evidence puts it')
        # With confidence defaulting to 1.0 the reason passes its full weight up, so the belief moves off 0.50.
        first = [ln for ln in r.stdout.splitlines() if ln.strip().startswith('1 ')]
        self.assertTrue(first, f'no line for page 1 in:\n{r.stdout}')
        self.assertNotIn('0.50', first[0], 'the belief did not move, so the default gate is not being used')
