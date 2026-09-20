# Build the workbook: the government ethics example (page 1, every sub-page, every specialized multiplier page) + blank templates
# (one per page kind), then verify every page against an independent model after a LibreOffice recalculation.
import json, os, shutil, subprocess, sys
from openpyxl import Workbook, load_workbook
HERE = os.path.dirname(os.path.abspath(__file__)); sys.path.insert(0, HERE)
import build_pages as bp
from build_pages import Page, CONSTS, is_page
from build_subpages import SubPage, KINDS

# Content comes either from the Python spec module or, preferably, from the flat data-entry workbook. The entry
# workbook is the authoring surface: two sheets of rows, no formatting, pages referred to by key rather than by
# tab number. Everything the reader sees is generated from it.
#   python3 build_example.py                       build from example_gov.py
#   python3 build_example.py ISE_Data_Entry.xlsx   build from the data-entry workbook
ENTRY = sys.argv[1] if len(sys.argv) > 1 and sys.argv[1].endswith('.xlsx') else None
if ENTRY:
    from ise_tables import read_entry, tables_to_specs
    SPECS, BELIEFS = tables_to_specs(*read_entry(ENTRY))
    print('loaded', len(SPECS), 'pages from', ENTRY)
else:
    import example_gov as _src
    SPECS, BELIEFS = _src.PAGES, getattr(_src, 'BELIEFS', {1})
class _Ez: pass
ez = _Ez(); ez.PAGES = SPECS

OUT = os.path.join(HERE, 'ISE_Government_Ethics_Example.xlsx')
bp.KIND_OF.clear(); bp.KIND_OF.update({pid: sp.get('kind', 'belief') for pid, sp in ez.PAGES.items()})

# Confidence is the one quantity the workbook does not compute itself: it needs the whole tree, so confidence.py
# computes it and every page carries its own value in a labelled cell, mirrored at W3 for parent rows to read.
# This is what keeps the workbook numerically identical to the website.
import render_site as _rs
_corpus = _rs.Corpus(ENTRY, 'workbook') if ENTRY else None
if _corpus is not None:
    bp.CONF_OF.clear(); bp.CONF_OF.update({pid: round(_corpus.conf.of(pid), 6) for pid in SPECS})
# Where each page's truth starts, and how heavily, from what the page says it rests on. Typed page data, so the
# workbook could compute it, but it is derived once here so the two never drift apart.
import evidence as _ev
bp.BASIS_OF.clear()
bp.BASIS_OF.update({pid: (round(_ev.prior(sp, 1)['p0'], 6), round(_ev.prior(sp, 1)['weight'], 6)) for pid, sp in SPECS.items()})
for _pid, _sp in SPECS.items(): _sp['_tab'] = _pid
wb = Workbook(); ws = wb.active; ws.title = 'Template'
tpl = Page(ws, None); tpl.build(); ws.sheet_properties.tabColor = '7F7F7F'
TPL = {'belief': tpl}
for kind in KINDS:
    w = wb.create_sheet(f'Template {KINDS[kind]["label"]}'); t = SubPage(w, {'kind': kind}); t.build(); TPL[kind] = t; w.sheet_properties.tabColor = 'A6A6A6'
TABCOLOR = {'belief': '1F3864', 'claim': '8EA9C1', 'linkage': '2E6F40', 'importance': 'D99B00', 'interest': 'F2C14E', 'uniqueness': '6A4C93', 'equivalence': '9A7FC1', 'driver': '2A9D8F', 'media': 'C0504D'}
KINDNAME = {'belief': 'Belief', 'claim': 'Claim', 'linkage': 'Linkage', 'importance': 'Importance', 'interest': 'Interest', 'uniqueness': 'Uniqueness', 'equivalence': 'Equivalence', 'driver': 'Driver', 'media': 'Media'}
def KIND(pid):
    sp = ez.PAGES[pid]
    return sp.get('kind') or ('belief' if pid in BELIEFS else 'claim')
PAGES = {}
for pid in sorted(ez.PAGES):
    sp = ez.PAGES[pid]; w = wb.create_sheet(str(pid))
    pg = SubPage(w, sp) if sp.get('kind') else Page(w, sp)
    pg.build(); PAGES[pid] = pg
    w.sheet_properties.tabColor = TABCOLOR[KIND(pid)]

# ---------------------------------------------------------------- Contents tab: every page, its question, score and completeness
from openpyxl.styles import Font, PatternFill, Alignment
from build_pages import TCELL, PAGE, HL, COMPLETE_CELL, LINK, GREY, HDR_FILL, WRAP, MIDC
ct = wb.create_sheet('Contents')
for col, wdt in (('A', 7), ('B', 13), ('C', 95), ('D', 8), ('E', 10), ('F', 60)): ct.column_dimensions[col].width = wdt
ct.sheet_view.showGridLines = False; ct.freeze_panes = 'A4'
ct['A1'] = 'Contents'; ct['A1'].font = Font(bold=True, size=14, color='FFFFFF'); ct['A1'].fill = HDR_FILL; ct.merge_cells('A1:F1')
n = len(ez.PAGES); first, last = 4, 4 + n - 1
ct['A2'] = f'=COUNTA(A{first}:A{last})&" pages.  "&COUNTIF(E{first}:E{last},"no")&" fail their completeness gate (one-sided or empty).  Every number is read from the page it names; click a tab number or a question to open it. The tab number is the page\'s ID (hidden on the web page)."'
ct['A2'].alignment = WRAP; ct.merge_cells('A2:F2'); ct.row_dimensions[2].height = 30
for col, h in (('A', 'Tab'), ('B', 'Kind'), ('C', 'Question or claim (read from the page)'), ('D', 'Score'), ('E', 'Complete'), ('F', 'Used on')):
    c = ct[f'{col}3']; c.value = h; c.font = Font(bold=True, size=9); c.fill = GREY; c.alignment = MIDC
for i, pid in enumerate(sorted(ez.PAGES)):
    r = first + i; sp = ez.PAGES[pid]; kind = KIND(pid)
    ct[f'A{r}'] = f'=HYPERLINK("#{PAGE(pid)}A1",{pid})'; ct[f'A{r}'].font = LINK; ct[f'A{r}'].alignment = MIDC
    ct[f'B{r}'] = KINDNAME[kind]; ct[f'B{r}'].font = Font(size=9); ct[f'B{r}'].alignment = MIDC
    ct[f'C{r}'] = f'={HL(pid, PAGE(pid) + "$C$1")}'; ct[f'C{r}'].font = LINK; ct[f'C{r}'].alignment = WRAP
    ct[f'D{r}'] = f'={HL(pid, TCELL(pid))}'; ct[f'D{r}'].number_format = '0.00'; ct[f'D{r}'].font = LINK; ct[f'D{r}'].alignment = MIDC
    ct[f'E{r}'] = f'=IF({PAGE(pid)}{COMPLETE_CELL}=1,"yes","no")'; ct[f'E{r}'].alignment = MIDC
    par = sp.get('supports')
    if par: ct[f'F{r}'] = f'={HL(par, PAGE(par) + "$C$1")}'; ct[f'F{r}'].font = Font(size=9, color='0563C1'); ct[f'F{r}'].alignment = WRAP
    ct.row_dimensions[r].height = 30 if kind in ('belief', 'claim', 'interest', 'media') else 58
    ct[f'A{r}'].fill = PatternFill('solid', fgColor=TABCOLOR[kind]) if kind != 'claim' else PatternFill('solid', fgColor='E8EEF4')
r = last + 2
ct[f'A{r}'] = 'Templates'; ct[f'A{r}'].font = Font(bold=True, size=10); r += 1
for name in ['Template'] + [f'Template {KINDS[k]["label"]}' for k in KINDS]:
    ct[f'A{r}'] = f'=HYPERLINK("#\'{name}\'!A1","{name}")'; ct[f'A{r}'].font = LINK; ct.merge_cells(f'A{r}:C{r}'); r += 1
ct.sheet_properties.tabColor = '000000'
wb._sheets = [ct] + [wb[str(pid)] for pid in sorted(ez.PAGES)] + [wb['Template']] + [wb[f'Template {KINDS[k]["label"]}'] for k in KINDS]
wb.active = 0
wb.calculation.fullCalcOnLoad = True
wb.save(OUT)
json.dump({'belief': tpl.rows, **{k: TPL[k].rows for k in KINDS}}, open(OUT.replace('.xlsx', '_rows.json'), 'w'), default=str)
print('saved', OUT, 'sheets', len(wb.sheetnames), 'layouts', {k: (v['truth'], v['last']) for k, v in bp.LAYOUTS.items()})

# ---------------------------------------------------------------- recalc
shutil.rmtree(os.path.join(HERE, 'ex_out'), ignore_errors=True)
subprocess.run(['soffice', '--headless', '--calc', '--convert-to', 'xlsx', '--outdir', os.path.join(HERE, 'ex_out'), OUT], capture_output=True, timeout=900)
RC = os.path.join(HERE, 'ex_out', os.path.basename(OUT))
wb2 = load_workbook(RC, data_only=True)

# ---------------------------------------------------------------- independent model (every multiplier is a page truth or a constant)
from score_reference import Model
model = Model(ez.PAGES, {k: v for k, _, v, _ in CONSTS})
if _corpus is not None: model.conf = _corpus.conf.of   # the workbook gates on confidence, so the check must too
def T(pid): return model.truth(pid)

bad = 0
for pid in sorted(ez.PAGES):
    m = model.evaluate(pid); w = wb2[str(pid)]; ENG = PAGES[pid].rows['engine']; kind = ez.PAGES[pid].get('kind', 'belief')
    checks = {'W1': m['truth'], 'X1': m['belief'], f'D{ENG["TRUTH"]}': m['truth'], f'D{ENG["BELIEF"]}': m['belief']}
    if kind != 'importance': checks.update({f'D{ENG["PRO"]}': m['pro'], f'D{ENG["CON"]}': m['con']})
    if kind == 'media': checks[f'D{ENG["IMPACT"]}'] = m['impact']
    if kind == 'belief':
        checks.update({f'D{ENG["SUPP"]}': m['supp'], f'D{ENG["WEAK"]}': m['weak'], f'D{ENG["PREDSCORE"]}': m['pred'], f'D{ENG["RAWTRUTH"]}': m['raw']})
    else:
        sp = ez.PAGES[pid]
        if is_page(sp.get('x')): checks['K4'] = T(sp['x'])
        yk = sp.get('y') if kind != 'uniqueness' else sp.get('z')
        if is_page(yk): checks['K5'] = T(yk)
    for ref, v in checks.items():
        got = w[ref].value
        if got is None or abs(float(got) - v) > 1e-6: bad += 1; print('MISMATCH', pid, kind, ref, got, v)
print('pages', len(ez.PAGES), 'mismatches', bad)
# audit: every interest listed on an importance page of a page-1 row is in page 1's interests table, and every page-1 interest is used somewhere
p1int = [r['id'] for r in ez.PAGES[1]['int_sup'] + ez.PAGES[1]['int_opp'] if r.get('id')]
used = {}
for lst in (ez.PAGES[1]['args']['agree'], ez.PAGES[1]['args']['disagree'], ez.PAGES[1]['evid']['for'], ez.PAGES[1]['evid']['against'], ez.PAGES[1]['pred_true'], ez.PAGES[1]['pred_false'], ez.PAGES[1]['media_for'], ez.PAGES[1]['media_against']):
    for row in lst:
        ip = row.get('imp')
        if is_page(ip):
            for d in ez.PAGES[ip].get('interests', []): used.setdefault(d['id'], []).append(ip)
for row in ez.PAGES[1]['benefits'] + ez.PAGES[1]['costs']:
    if is_page(row.get('who')): used.setdefault(row['who'], []).append('cba')
print('audit: interests on importance pages not in the page-1 table:', sorted(set(used) - set(p1int)), '| page-1 interests no row speaks to:', sorted(set(p1int) - set(used)))
c1 = wb2['1']; E1 = PAGES[1].rows['engine']
print('page 1 coverage/completeness:', {k: c1[f'D{E1[k]}'].value for k in ('NROWS', 'NOLINK', 'NOIMP', 'NOUNIQ', 'COMPLETE', 'NCHILD', 'NINCOMPLETE', 'INCOMPLETE', 'RAWTRUTH', 'TRUTH', 'WEAKEST')})
inc = [pid for pid in sorted(ez.PAGES) if wb2[str(pid)]['W2'].value == 0]
print('pages failing completeness:', len(inc), inc[:30])
print('contents A2:', wb2['Contents']['A2'].value)
errs = [(s.title, c.coordinate, c.value) for s in wb2.worksheets for row in s.iter_rows() for c in row if isinstance(c.value, str) and ((c.value.startswith('#') and c.value != '#') or c.value.startswith('Err:'))]
print('formula errors', len(errs), errs[:8])
emd = [(s.title, c.coordinate) for s in wb2.worksheets for row in s.iter_rows() for c in row if isinstance(c.value, str) and ('—' in c.value or 'utiliz' in c.value.lower())]
print('em dashes / utilize', len(emd), emd[:5])
r1 = wb2['1']
print('--- page 1 header:', r1['C1'].value)
print('--- formula-built questions (one per kind and type)')
shown = set()
for p in sorted(ez.PAGES):
    k = ez.PAGES[p].get('kind'); tag = (k, ez.PAGES[p].get('typ'))
    if k and k not in ('interest', 'media') and tag not in shown:
        shown.add(tag); print(p, tag, '|', repr(wb2[str(p)]['C1'].value))
imp2 = ez.PAGES[1]['args']['agree'][0]['imp']
print(f'--- page {imp2} (importance of reason 2) readout and interests')
w = wb2[str(imp2)]; print(w[f'B{PAGES[imp2].rows["readout"]}'].value)
for r in PAGES[imp2].rows['interests']: print([w[f'{c}{r}'].value for c in 'ABCFGHIJLS'])
ints = [p for p in sorted(ez.PAGES) if ez.PAGES[p].get('kind') == 'interest']
med = [p for p in sorted(ez.PAGES) if ez.PAGES[p].get('kind') == 'media']
i0 = ints[0]
print('--- interest page', i0, ':', wb2[str(i0)]['C1'].value, '| D4:', wb2[str(i0)]['D4'].value, '| F2:', wb2[str(i0)]['F2'].value)
print('    readout:', str(wb2[str(i0)][f'B{PAGES[i0].rows["readout"]}'].value)[:150])
print('--- media page', med[0], ':', wb2[str(med[0])]['C1'].value, '|', wb2[str(med[0])][f'B{PAGES[med[0]].rows["readout"]}'].value)
for b in (1, 2, 3, 4, 5):
    w = wb2[str(b)]; E = PAGES[b].rows['engine']
    tr = w['D' + str(E['TRUTH'])].value; raw = w['D' + str(E['RAWTRUTH'])].value if 'RAWTRUTH' in E else None
    print('--- page', b, 'truth', tr, 'argued', raw, '|', str(w['C1'].value)[:80])
print('templates:', [s for s in wb2.sheetnames if s.startswith('Template')])
tl = wb2['Template Linkage']; print('template linkage C1:', tl['C1'].value)

# ---------------------------------------------------------------- database export (same normalized tables the model scored)
from export_db import export
pages_, edges_ = export(ez.PAGES, {k: v for k, _, v, _ in CONSTS}, os.path.join(HERE, 'db'), stem='ise_gov_ethics', const_meanings={k: m for k, _, _, m in CONSTS}, beliefs=BELIEFS)
print('db export: pages', len(pages_), 'edges', len(edges_))
